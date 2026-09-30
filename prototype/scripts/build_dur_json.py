"""프로토타입용 dur.json 만들기 (2026-09-28 형준)

원본: 심평원 약제급여목록(내복) · DUR m.db · 효능군중복 품목리스트
사용: python3 scripts/build_dur_json.py --raw <data_raw 폴더> --out src/data/dur.json
기준일(--today)에 적용 중인 규칙만 넣음. 기본값 20260928
"""

import argparse
import json
import os
import sqlite3

import openpyxl

ap = argparse.ArgumentParser()
ap.add_argument("--raw", required=True, help="data_raw 폴더 (dur/ · hira/ 가 있는 곳)")
ap.add_argument("--out", required=True)
ap.add_argument("--today", default="20260928")
args = ap.parse_args()
today = args.today
raw = args.raw

drug_list = os.path.join(raw, "hira", "약제급여목록및급여상한금액표_(2026.9.1.)_공개용(비인가자) 1부.xlsx")
wb = openpyxl.load_workbook(drug_list, read_only=True)
prods = []
seen = set()
for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
    if not r[8] or r[1] != "내복":
        continue
    edi = str(r[8]).zfill(9)
    if edi in seen:
        continue
    seen.add(edi)
    # [제품코드, 제품명, 주성분코드, 성분명, 업체, 전문/일반] + 아래에서 식약 분류번호
    prods.append(
        [
            edi,
            str(r[9]).strip(),
            str(r[5]).strip(),
            str(r[7] or "").strip()[:60],
            str(r[10] or "").strip(),
            str(r[14] or ""),
        ]
    )
# 7번째 칸: 식약 분류번호 (같은 제품코드가 여러 행이면 마지막 행 값)
cls = {}
for r in wb.worksheets[0].iter_rows(min_row=2, values_only=True):
    if r[8]:
        cls[str(r[8]).zfill(9)] = str(r[3] or "")
for p in prods:
    p.append(cls.get(p[0], ""))
codes = {p[2] for p in prods}

db = sqlite3.connect(os.path.join(raw, "dur", "m.db"))
reasons = []
reason_index = {}


def reason_id(text):
    text = (text or "").strip()
    if text not in reason_index:
        reason_index[text] = len(reasons)
        reasons.append(text)
    return reason_index[text]


def active(fr, to):
    return fr <= today <= to


# 병용금기 성분 쌍 → [사유, 공고일]
inter = {}
sql = "select DUR_CD_A, DUR_CD_B, ADPT_FR_DT, ADPT_TO_DT, DUR_SD_EFT, ANNCE_DT from TBJBD43"
for code_a, code_b, fr, to, eft, annce_dt in db.execute(sql):
    if not active(fr, to) or code_a not in codes or code_b not in codes:
        continue
    x, y = sorted([code_a, code_b])
    inter.setdefault(x + "|" + y, [reason_id(eft), annce_dt or ""])

# 병용금기 예외
exc = {}
for row in db.execute("select * from TBJBD46"):
    code_a, code_b, fr, to, medc = row[:5]
    if active(fr, to):
        x, y = sorted([code_a, code_b])
        exc.setdefault(x + "|" + y, []).append(medc)

# 노인주의
eld = {}
for row in db.execute("select * from TBDUC230"):
    code, fr, to, age, txt = row[0], row[1], row[2], row[4], row[7]
    if active(fr, to) and code in codes:
        eld[code] = [age, reason_id(txt)]

# 연령금기
agel = {}
for row in db.execute("select * from TBJBD44"):
    code, fr, to, age, unit, eft, cond = row[0], row[1], row[2], row[3], row[4], row[6], row[11]
    if active(fr, to) and code in codes:
        agel[code] = [age, unit, cond, reason_id(eft)]

# 효능군중복
wb = openpyxl.load_workbook(os.path.join(raw, "dur", "게시_효능군중복 품목리스트_2609.xlsx"), read_only=True)
eff = {}
for ws in wb.worksheets:
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[5]:
            eff[str(r[5]).zfill(9)] = [str(r[4]), str(r[0])]
eff = {k: v for k, v in eff.items() if k in seen}

meta = {
    "source": "심평원 약제급여목록(2026-09-01, 내복) · DUR m.db(2026-08-24) · 효능군중복 품목리스트(2026-09)",
    "built": today,
}
os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
out = {
    "meta": meta,
    "products": prods,
    "reasons": reasons,
    "interactions": inter,
    "exceptions": exc,
    "elderly": eld,
    "ageLimit": agel,
    "efficacy": eff,
}
with open(args.out, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
size = round(os.path.getsize(args.out) / 1e6, 2)
print(len(prods), "products", len(inter), "pairs", len(eld), "elderly", len(agel), "age", len(eff), "eff", size, "MB")
